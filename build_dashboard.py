"""Gera dashboard_xbox.xlsx a partir de base.xlsx (abas Cálculos e Dashboard)."""
import openpyxl
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("base.xlsx")
ws_assets, ws_base, ws_calc, ws_dash = wb.worksheets

GREEN, LIME, MINT, LIGHT = "22C55E", "9BC848", "2AE6B1", "5BF6A8"
BG, CARD, DARK = "E8E6E9", "FFFFFF", "1A1A1A"

# Tradução da base (cabeçalhos e valores) para português
HEADERS = ["ID do Assinante", "Nome", "Plano", "Data de Início", "Renovação Automática", "Preço da Assinatura",
           "Tipo de Assinatura", "Passe de Temporada EA Play", "Preço do Passe EA Play",
           "Passe de Temporada Minecraft", "Preço do Passe Minecraft", "Valor do Cupom", "Valor Total"]
VALUES = {"Yes": "Sim", "No": "Não", "Monthly": "Mensal", "Quarterly": "Trimestral", "Annual": "Anual"}
for j, h in enumerate(HEADERS, 1):
    ws_base.cell(1, j, h)
tbl = ws_base.tables["Tabela1"]
for col, h in zip(tbl.tableColumns, HEADERS):
    col.name = h
for row in ws_base.iter_rows(min_row=2):
    for cell in row:
        if cell.value in VALUES:
            cell.value = VALUES[cell.value]

n = ws_base.max_row  # última linha de dados
B = f"'{ws_base.title}'!"


def rng(col):
    return f"{B}${col}$2:${col}${n}"


ID, PLAN, DATE, RENEW, PRICE, TYPE = rng("A"), rng("C"), rng("D"), rng("E"), rng("F"), rng("G")
EA, EAP, MC, MCP, COUP, TOTAL = rng("H"), rng("I"), rng("J"), rng("K"), rng("L"), rng("M")

# ---------- Cálculos ----------
c = ws_calc
c.sheet_view.showGridLines = False
head = Font(bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=DARK)


def table(r, c0, title, cols, rows):
    """Escreve um bloco de cálculo; devolve (linha_primeira, linha_última)."""
    c.cell(r, c0, title).font = Font(bold=True, size=12)
    for j, h in enumerate(cols):
        x = c.cell(r + 1, c0 + j, h)
        x.font, x.fill = head, head_fill
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c.cell(r + 2 + i, c0 + j, v)
    return r + 2, r + 1 + len(rows)


# KPIs
c["A1"] = "Indicadores"
c["A1"].font = Font(bold=True, size=14)
kpis = [
    ("Assinantes", f"=COUNTA({ID})"),
    ("Receita total", f"=SUM({TOTAL})"),
    ("Ticket médio", "=B4/B3"),
    ("Total em cupons", f"=SUM({COUP})"),
    ("% renovação automática", f'=COUNTIF({RENEW},"Sim")/B3'),
]
for i, (k, f) in enumerate(kpis):
    c.cell(3 + i, 1, k)
    c.cell(3 + i, 2, f)
c["B4"].number_format = c["B5"].number_format = c["B6"].number_format = '"US$" #,##0.00'
c["B7"].number_format = "0.0%"

plans = ["Core", "Standard", "Ultimate"]
p0, p1 = table(10, 1, "Por plano", ["Plano", "Assinantes", "Receita"],
               [(p, f"=COUNTIF({PLAN},A{12+i})", f"=SUMIF({PLAN},A{12+i},{TOTAL})") for i, p in enumerate(plans)])

types = [("Mensal", "Mensal"), ("Trimestral", "Trimestral"), ("Anual", "Anual")]
t0, t1 = table(17, 1, "Por tipo de assinatura", ["Tipo", "Assinantes", "Receita", "Chave"],
               [(lbl, f"=COUNTIF({TYPE},D{19+i})", f"=SUMIF({TYPE},D{19+i},{TOTAL})", k) for i, (k, lbl) in enumerate(types)])

r0, r1 = table(24, 1, "Renovação automática", ["Renovação", "Assinantes", "Chave"],
               [("Com renovação", f'=COUNTIF({RENEW},C26)', "Sim"), ("Sem renovação", f'=COUNTIF({RENEW},C27)', "Não")])

a0, a1 = table(30, 1, "Passes adicionais", ["Passe", "Assinantes", "Receita"],
               [("EA Play", f'=COUNTIF({EA},"Sim")', f"=SUM({EAP})"),
                ("Minecraft", f'=COUNTIF({MC},"Sim")', f"=SUM({MCP})")])

months = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
m_rows = []
for m in range(12):
    r = 36 + m
    start = f"DATE(2024,{m+1},1)"
    end = f"DATE(2024,{m+2},1)" if m < 11 else "DATE(2025,1,1)"
    m_rows.append((months[m], f'=COUNTIFS({DATE},">="&{start},{DATE},"<"&{end})',
                   f'=SUMIFS({TOTAL},{DATE},">="&{start},{DATE},"<"&{end})'))
m0, m1 = table(33, 1, "Por mês de início (2024)", ["Mês", "Novas assinaturas", "Receita"], m_rows)

for col, w in zip("ABCD", (26, 18, 16, 12)):
    c.column_dimensions[col].width = w
for r in list(range(p0, p1 + 1)) + list(range(t0, t1 + 1)) + list(range(a0, a1 + 1)) + list(range(m0, m1 + 1)):
    c.cell(r, 3).number_format = '"US$" #,##0.00'

# ---------- Dashboard ----------
d = ws_dash
d.sheet_view.showGridLines = False
d.sheet_view.zoomScale = 90
NCOL = 24
for j in range(1, NCOL + 1):
    d.column_dimensions[get_column_letter(j)].width = 6.5
for r in range(1, 60):
    for j in range(1, NCOL + 1):
        d.cell(r, j).fill = PatternFill("solid", fgColor=BG)

for r in range(1, 5):
    for j in range(1, NCOL + 1):
        d.cell(r, j).fill = PatternFill("solid", fgColor=DARK)
d["B2"] = "Xbox Game Pass | Dashboard de Vendas 2024"
d["B2"].font = Font(bold=True, size=22, color=LIME)
d["B3"] = "Assinantes, receita e passes adicionais"
d["B3"].font = Font(size=11, color="FFFFFF")
d.row_dimensions[2].height = 32

calc = f"'{c.title}'!"
cards = [  # (coluna inicial, rótulo, fórmula, formato)
    (2, "Assinantes", f"={calc}B3", "#,##0"),
    (6, "Receita total", f"={calc}B4", '"US$" #,##0'),
    (10, "Ticket médio", f"={calc}B5", '"US$" #,##0.00'),
    (14, "Total em cupons", f"={calc}B6", '"US$" #,##0'),
    (18, "Renovação automática", f"={calc}B7", "0.0%"),
]
for col, label, f, fmt in cards:
    for r in (6, 7, 8):
        for j in range(col, col + 4):
            d.cell(r, j).fill = PatternFill("solid", fgColor=CARD)
    d.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 3)
    d.merge_cells(start_row=7, start_column=col, end_row=8, end_column=col + 3)
    d.cell(6, col, label).font = Font(size=10, color="595959")
    v = d.cell(7, col, f)
    v.number_format, v.font = fmt, Font(bold=True, size=22, color=DARK)
    d.cell(6, col).alignment = d.cell(7, col).alignment = Alignment(horizontal="center", vertical="center")
    for j in range(col, col + 4):
        d.cell(5, j).fill = PatternFill("solid", fgColor=GREEN)
    d.row_dimensions[5].height = 4

COLORS = [GREEN, LIME, MINT, LIGHT]


def style(ch, title, w=11.8, h=7.6):
    ch.title, ch.width, ch.height = title, w, h
    ch.legend = None
    return ch


def colors(series, pts=None):
    series.graphicalProperties.solidFill = GREEN
    series.graphicalProperties.line.solidFill = GREEN
    if pts:
        for i in range(pts):
            p = DataPoint(idx=i)
            p.graphicalProperties.solidFill = COLORS[i % 4]
            series.dPt.append(p)


def labels(series):
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    for a in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(series.dLbls, a, False)


def ref(col, r0_, r1_):
    return Reference(c, min_col=col, min_row=r0_, max_row=r1_)


# 1) Receita por plano
ch = style(BarChart(), "Receita por plano (US$)")
ch.type = "col"
ch.add_data(ref(3, p0 - 1, p1), titles_from_data=True)
ch.set_categories(ref(1, p0, p1))
colors(ch.series[0], 3), labels(ch.series[0])
ch.y_axis.delete = ch.x_axis.delete = False
d.add_chart(ch, "B10")

# 2) Assinantes por tipo
ch = style(DoughnutChart(), "Assinantes por tipo de assinatura")
ch.add_data(ref(2, t0 - 1, t1), titles_from_data=True)
ch.set_categories(ref(1, t0, t1))
colors(ch.series[0], 3), labels(ch.series[0])
ch.legend = openpyxl.chart.legend.Legend()
ch.legend.position = "b"
d.add_chart(ch, "J10")

# 3) Renovação automática
ch = style(DoughnutChart(), "Renovação automática")
ch.add_data(ref(2, r0 - 1, r1), titles_from_data=True)
ch.set_categories(ref(1, r0, r1))
colors(ch.series[0], 2), labels(ch.series[0])
ch.legend = openpyxl.chart.legend.Legend()
ch.legend.position = "b"
d.add_chart(ch, "R10")

# 4) Novas assinaturas por mês
ch = style(LineChart(), "Novas assinaturas por mês", w=18, h=7.6)
ch.add_data(ref(2, m0 - 1, m1), titles_from_data=True)
ch.set_categories(ref(1, m0, m1))
ch.series[0].graphicalProperties.line.solidFill = GREEN
ch.series[0].graphicalProperties.line.width = 28000
ch.series[0].marker.symbol = "circle"
ch.series[0].smooth = False
labels(ch.series[0])
ch.y_axis.delete = ch.x_axis.delete = False
d.add_chart(ch, "B26")

# 5) Passes adicionais
ch = style(BarChart(), "Receita dos passes adicionais (US$)", w=11.8)
ch.type = "bar"
ch.add_data(ref(3, a0 - 1, a1), titles_from_data=True)
ch.set_categories(ref(1, a0, a1))
colors(ch.series[0], 2), labels(ch.series[0])
ch.y_axis.delete = ch.x_axis.delete = False
d.add_chart(ch, "R26")

# 6) Receita por tipo de assinatura
ch = style(BarChart(), "Receita por tipo de assinatura (US$)", w=18)
ch.type = "col"
ch.add_data(ref(3, t0 - 1, t1), titles_from_data=True)
ch.set_categories(ref(1, t0, t1))
colors(ch.series[0], 3), labels(ch.series[0])
ch.y_axis.delete = ch.x_axis.delete = False
d.add_chart(ch, "B42")

wb.active = wb.worksheets.index(d)
for ws in wb:
    ws.sheet_view.tabSelected = ws is d
wb.save("dashboard_xbox.xlsx")
print("ok")
